# content_extractor.py
import os
import hashlib
import pypdf
import docx
import pptx
import openpyxl

def get_file_metadata(file_path):
    if not os.path.exists(file_path):
        return None
    stat = os.stat(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    
    # Compute MD5
    hasher = hashlib.md5()
    try:
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(65536), b''):
                hasher.update(chunk)
        file_hash = hasher.hexdigest()
    except Exception:
        file_hash = 'unknown'

    return {
        'name': os.path.basename(file_path),
        'path': os.path.abspath(file_path),
        'size': stat.st_size,
        'modified': stat.st_mtime,
        'extension': ext,
        'md5': file_hash
    }

def extract_content(file_path, max_chars=2000):
    if not os.path.exists(file_path):
        return ''
    ext = os.path.splitext(file_path)[1].lower()
    text = ''
    try:
        if ext == '.pdf':
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages[:5]:
                t = page.extract_text()
                if t:
                    text += t + '\n'
        elif ext == '.docx':
            doc = docx.Document(file_path)
            for p in doc.paragraphs[:30]:
                if p.text:
                    text += p.text + '\n'
        elif ext == '.pptx':
            prs = pptx.Presentation(file_path)
            for slide in prs.slides[:5]:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text += shape.text_frame.text + '\n'
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
            sheet = wb.active
            rows = []
            for row in list(sheet.iter_rows(max_row=8, values_only=True)):
                rows.append(' | '.join([str(c) for c in row if c is not None]))
            text = '\n'.join(rows)
        elif ext in ['.txt', '.md', '.html', '.css', '.js', '.py', '.c', '.cpp', '.sql', '.json', '.csv', '.xml', '.yml', '.yaml', '.ini', '.log']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read(max_chars * 2)
        else:
            text = f'[{ext} binary file]'
    except Exception as e:
        text = f'[Error extracting text: {e}]'

    # Clean whitespace
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return '\n'.join(lines)[:max_chars]
